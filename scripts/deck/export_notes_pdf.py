"""Export speaker notes from the deck to a clean text-only PDF.

Reads:  MedQuAD_Robustness_Final_Presentation_v2_with_notes.pptx
        (falls back to _n50_polished_with_notes.pptx)
Writes: artifacts/n50_speaker_notes.pdf

The PDF is text-only — no slide thumbnails. Each slide gets:
  • a header row (slide number + slide title)
  • the speaker notes body, formatted with paragraph breaks
"""

from __future__ import annotations

import re
from pathlib import Path

from pptx import Presentation
from pptx.util import Emu
from reportlab.lib.colors import HexColor
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import LETTER
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import (
    KeepTogether,
    PageBreak,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
)

REPO = Path(__file__).resolve().parents[2]

_CANDIDATE_NAMES = (
    "MedQuAD_Robustness_Final_Presentation_v2_with_notes.pptx",
    "MedQuAD_Robustness_Final_Presentation_n50_polished_with_notes.pptx",
)
SRC = next(
    (REPO / n for n in _CANDIDATE_NAMES if (REPO / n).exists()),
    REPO / _CANDIDATE_NAMES[0],
)
OUT = REPO / "artifacts" / "n50_speaker_notes.pdf"

# Palette echoes the deck for visual continuity
TEAL = HexColor("#127C7A")
SLATE = HexColor("#25323D")
GOLD = HexColor("#E0A13A")
MUTED = HexColor("#52616B")


SECTION_RE = re.compile(r"^\s*\d{2}\s*\|\s*.+$")


# Per-slide title overrides for slides where shape-scraping doesn't pick up
# the right thing (slide 1 is a banner-style title; slide 19 has no header).
_TITLE_OVERRIDES: dict[int, str] = {
    1: "Title — Medical LLM Robustness",
    19: "Thank you",
}


def _slide_title(slide, slide_idx: int) -> str:
    """Pick the most-likely-title text from a slide.

    The polished deck splits the section header — '01' lives in a small motif
    circle, the title word ('Problem') in a separate frame. We grab the first
    decent text after any motif markers as the title, then append the
    headline below it.
    """
    if slide_idx in _TITLE_OVERRIDES:
        return _TITLE_OVERRIDES[slide_idx]

    candidates = []
    for shape in slide.shapes:
        if not shape.has_text_frame:
            continue
        if shape.top is None:
            continue
        top_in = Emu(shape.top).inches
        if top_in > 1.5:
            continue  # below the header band
        txt = shape.text_frame.text.strip()
        if not txt or txt in {"01", "02", "03", "04", "05", "06", "07", "08", "09", "10", "11", "12", "13", "14", "15", "16", "17"}:
            continue
        if SECTION_RE.match(txt):
            txt = txt.split("|", 1)[1].strip()
        candidates.append((top_in, txt))

    candidates.sort()
    if not candidates:
        return ""

    parts = []
    for _, t in candidates[:2]:
        if t and t not in parts:
            parts.append(t)
    return " — ".join(parts)


def _notes(slide) -> str:
    if not slide.has_notes_slide:
        return ""
    return slide.notes_slide.notes_text_frame.text


def main() -> None:
    if not SRC.exists():
        raise FileNotFoundError(SRC)
    OUT.parent.mkdir(parents=True, exist_ok=True)

    prs = Presentation(SRC)
    slide_count = len(prs.slides)

    # --- styles -------------------------------------------------------------
    base = getSampleStyleSheet()
    deck_title_style = ParagraphStyle(
        "DeckTitle",
        parent=base["Title"],
        fontName="Helvetica-Bold",
        fontSize=22,
        leading=28,
        textColor=SLATE,
        spaceAfter=4,
        alignment=TA_LEFT,
    )
    deck_subtitle_style = ParagraphStyle(
        "DeckSubtitle",
        parent=base["Normal"],
        fontName="Helvetica",
        fontSize=11,
        leading=14,
        textColor=MUTED,
        spaceAfter=18,
    )
    slide_num_style = ParagraphStyle(
        "SlideNum",
        parent=base["Normal"],
        fontName="Helvetica-Bold",
        fontSize=10,
        leading=12,
        textColor=TEAL,
        spaceAfter=2,
    )
    slide_title_style = ParagraphStyle(
        "SlideTitle",
        parent=base["Heading2"],
        fontName="Helvetica-Bold",
        fontSize=15,
        leading=19,
        textColor=SLATE,
        spaceAfter=10,
    )
    body_style = ParagraphStyle(
        "Body",
        parent=base["Normal"],
        fontName="Helvetica",
        fontSize=11,
        leading=15,
        textColor=SLATE,
        spaceAfter=8,
    )
    cue_style = ParagraphStyle(
        "Cue",
        parent=body_style,
        fontName="Helvetica-Oblique",
        textColor=GOLD,
        spaceBefore=4,
    )

    # --- build doc ----------------------------------------------------------
    story: list = []

    story.append(Paragraph("Medical LLM Robustness", deck_title_style))
    story.append(
        Paragraph(
            "Speaker notes &middot; n=50 final presentation &middot; Group 2",
            deck_subtitle_style,
        )
    )

    for idx, slide in enumerate(prs.slides, start=1):
        title = _slide_title(slide, idx)
        notes = _notes(slide).strip()

        block: list = []
        block.append(Paragraph(f"SLIDE {idx} OF {slide_count}", slide_num_style))
        if title:
            block.append(Paragraph(_escape(title), slide_title_style))

        if not notes:
            block.append(Paragraph("<i>(no notes)</i>", body_style))
        else:
            for paragraph in notes.split("\n"):
                p = paragraph.strip()
                if not p:
                    block.append(Spacer(1, 4))
                    continue
                style = cue_style if p.startswith(("→", "->")) else body_style
                block.append(Paragraph(_escape(p), style))

        story.extend(block)
        if idx < slide_count:
            story.append(Spacer(1, 12))
            story.append(PageBreak())

    doc = SimpleDocTemplate(
        str(OUT),
        pagesize=LETTER,
        leftMargin=0.85 * inch,
        rightMargin=0.85 * inch,
        topMargin=0.75 * inch,
        bottomMargin=0.75 * inch,
        title="Medical LLM Robustness — Speaker Notes",
        author="Group 2",
    )
    doc.build(story, onFirstPage=_footer, onLaterPages=_footer)
    print(f"Wrote {OUT}")


_BOLD_RE = re.compile(r"\*\*(.+?)\*\*")


def _escape(text: str) -> str:
    """HTML-escape, convert markdown **bold** -> <b>bold</b>, normalize arrow."""
    text = (
        text.replace("&", "&amp;")
        .replace("<", "&lt;")
        .replace(">", "&gt;")
        .replace("→", "&rarr;")
    )
    # Convert **bold** to ReportLab <b> tags. Run AFTER HTML escaping so the
    # bold tags themselves don't get escaped.
    text = _BOLD_RE.sub(r"<b>\1</b>", text)
    return text


def _footer(canvas, doc) -> None:
    canvas.saveState()
    canvas.setFont("Helvetica", 8)
    canvas.setFillColor(MUTED)
    canvas.drawString(
        0.85 * inch, 0.45 * inch, "MedQuAD Robustness  |  Group 2  |  n=50 speaker notes"
    )
    canvas.drawRightString(
        LETTER[0] - 0.85 * inch, 0.45 * inch, f"Page {doc.page}"
    )
    canvas.restoreState()


if __name__ == "__main__":
    main()
